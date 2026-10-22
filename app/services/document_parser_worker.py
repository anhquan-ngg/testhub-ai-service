"""Disposable parser process. No network, DB, provider calls or source logging."""

import hashlib
from importlib.metadata import version
import io
import json
import os
import sys
import zipfile
from pathlib import Path

from app.schemas.document import (
    ExtractionResult, MAX_FILE_BYTES, MAX_RESULT_BYTES, MAX_TEXT_CHARS,
    MAX_UNITS, SourceSpan,
)


class DocumentRejected(Exception):
    """Only an allowlisted error code may cross the process boundary."""


def checked_pptx(data: bytes) -> io.BytesIO:
    """Repack only bytes actually inflated within a cumulative expansion limit."""
    target = io.BytesIO()
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as source, zipfile.ZipFile(target, "w") as dest:
            entries = source.infolist()
            if len(entries) > 5000:
                raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")
            names: set[str] = set()
            total = 0
            for entry in entries:
                name = entry.filename
                if (name in names or name.startswith("/") or "\\" in name
                    or any(part in (".", "..") for part in name.split("/"))
                    or entry.flag_bits & 1 or entry.compress_type not in (0, 8)):
                    raise DocumentRejected("DOCUMENT_INVALID")
                names.add(name)
                remaining = 100 * 1024 * 1024 - total
                if entry.file_size > remaining:
                    raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")
                # A forged central-directory size is not the expansion boundary:
                # read at most remaining+1, verify actual length, then repack.
                with source.open(entry) as stream:
                    content = stream.read(remaining + 1)
                total += len(content)
                if total > 100 * 1024 * 1024:
                    raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")
                dest.writestr(name, content)
            if not {"[Content_Types].xml", "ppt/presentation.xml"}.issubset(names):
                raise DocumentRejected("DOCUMENT_INVALID")
    except (zipfile.BadZipFile, RuntimeError, OSError, NotImplementedError):
        raise DocumentRejected("DOCUMENT_INVALID") from None
    target.seek(0)
    return target


def extract(path: Path, source_type: str, include_notes: bool) -> ExtractionResult:
    """Preserve page/slide identities without rendering or executing embedded data."""
    data = path.read_bytes()
    if not data or len(data) > MAX_FILE_BYTES:
        raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")
    spans: list[SourceSpan] = []
    needs_ocr = False
    chars = 0

    def add(kind: str, position: int, text: str) -> None:
        nonlocal chars
        text = text.replace("\x00", "").strip()
        chars += len(text)
        if chars > MAX_TEXT_CHARS:
            raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")
        spans.append(SourceSpan(kind=kind, position=position, text=text))

    if source_type == "PDF":
        from pypdf import PdfReader
        if not data.startswith(b"%PDF-"):
            raise DocumentRejected("DOCUMENT_INVALID")
        reader = PdfReader(io.BytesIO(data), strict=True)
        if reader.is_encrypted:
            raise DocumentRejected("DOCUMENT_ENCRYPTED")
        count = len(reader.pages)
        if not 1 <= count <= MAX_UNITS:
            raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")
        for number, page in enumerate(reader.pages, 1):
            text = page.extract_text() or ""
            add("PAGE", number, text)
            # Conservatively flag image-only pages even in a mixed PDF.
            if not text.strip() and bool(page.images):
                needs_ocr = True
    elif source_type == "PPTX":
        from pptx import Presentation
        presentation = Presentation(checked_pptx(data))
        count = len(presentation.slides)
        if not 1 <= count <= MAX_UNITS:
            raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")

        def shape_text(shapes: object) -> list[str]:
            parts: list[str] = []
            for shape in shapes:
                if shape.shape_type == 6:  # GROUP
                    parts.extend(shape_text(shape.shapes))
                elif shape.has_table:
                    parts.extend("\t".join(cell.text for cell in row.cells) for row in shape.table.rows)
                elif shape.has_text_frame:
                    parts.append(shape.text_frame.text)
            return parts

        def contains_image(shapes: object) -> bool:
            return any(shape.shape_type == 13 or
                       (shape.shape_type == 6 and contains_image(shape.shapes))
                       for shape in shapes)

        for number, slide in enumerate(presentation.slides, 1):
            text = "\n".join(shape_text(slide.shapes))
            add("SLIDE", number, text)
            if not text.strip() and contains_image(slide.shapes):
                needs_ocr = True
            if include_notes and slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None and frame.text.strip():
                    add("NOTES", number, frame.text)
    else:
        raise DocumentRejected("DOCUMENT_INVALID")
    if not any(span.text for span in spans):
        needs_ocr = True
    empty = any(not span.text for span in spans if span.kind != "NOTES")
    return ExtractionResult(
        source_type=source_type, source_sha256=hashlib.sha256(data).hexdigest(),
        parser_version=f"{source_type.lower()}/{version('pypdf' if source_type == 'PDF' else 'python-pptx')};text/v1",
        include_notes=include_notes, unit_count=count, spans=spans,
        status="OCR_REQUIRED" if needs_ocr else "EXTRACTED",
        warnings=(["EMPTY_UNITS"] if empty else []) + (["OCR_REQUIRED"] if needs_ocr else []),
    )


def main() -> None:
    """Apply POSIX memory/CPU caps before importing third-party parsers."""
    if os.name == "posix":
        import resource
        resource.setrlimit(resource.RLIMIT_AS, (768 * 1024 * 1024,) * 2)
        resource.setrlimit(resource.RLIMIT_CPU, (40, 40))
        resource.setrlimit(resource.RLIMIT_FSIZE, (MAX_RESULT_BYTES,) * 2)
    output = Path(sys.argv[2])
    try:
        result = extract(Path(sys.argv[1]), sys.argv[3], sys.argv[4] == "true")
        encoded = result.model_dump_json().encode("utf-8")
        if len(encoded) > MAX_RESULT_BYTES:
            raise DocumentRejected("DOCUMENT_LIMIT_EXCEEDED")
    except DocumentRejected as exc:
        encoded = json.dumps({"error": str(exc)}).encode()
    except MemoryError:
        encoded = b'{"error":"DOCUMENT_LIMIT_EXCEEDED"}'
    except Exception:
        encoded = b'{"error":"DOCUMENT_INVALID"}'
    output.write_bytes(encoded)


if __name__ == "__main__":
    main()
