"""Real PDF/PPTX extraction through the production route and disposable parser."""

import asyncio
import io
from pathlib import Path
import subprocess
import sys
import threading
import zipfile

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject
from pptx import Presentation

from app.config import Settings
from app.main import create_app
from app.services.document_parser import DocumentParser, ExtractionError

TOKEN = "document-test-token-" + "a" * 32


def pdf_bytes(*, blank=False, encrypted=False):
    """Build a tiny text-layer PDF without adding fixture-generation dependencies."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=600, height=800)
    if not blank:
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
        stream = DecodedStreamObject()
        stream.set_data(b"BT /F1 12 Tf 30 700 Td (Stack uses LIFO.) Tj ET")
        page[NameObject("/Contents")] = writer._add_object(stream)
    if encrypted:
        writer.encrypt("test-only-password")
    target = io.BytesIO()
    writer.write(target)
    return target.getvalue()


def pptx_bytes():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Ngăn xếp"
    slide.placeholders[1].text = "Nguyên tắc LIFO"
    slide.notes_slide.notes_text_frame.text = "Ghi chú dành cho giảng viên"
    target = io.BytesIO()
    presentation.save(target)
    return target.getvalue()


@pytest.fixture
def client():
    with TestClient(create_app(Settings(internal_api_token=TOKEN))) as value:
        yield value


def extract(client, data, source="PDF", notes=False, **kwargs):
    return client.post(f"/internal/documents/extract?source_type={source}&include_notes={str(notes).lower()}",
        content=data, headers={"X-Internal-Token": TOKEN, "Content-Type": "application/octet-stream"}, **kwargs)


def test_text_pdf_preserves_page_and_hash(client):
    import hashlib
    data = pdf_bytes()
    response = extract(client, data)
    assert response.status_code == 200
    result = response.json()
    assert result["status"] == "EXTRACTED"
    assert result["source_sha256"] == hashlib.sha256(data).hexdigest()
    assert result["spans"] == [{"kind": "PAGE", "position": 1, "text": "Stack uses LIFO."}]


def test_exported_contract_matches_production_route():
    import json
    from pathlib import Path
    from scripts.export_document_contract import document_contract
    path = Path(__file__).resolve().parents[1] / "contracts/document-extraction.openapi.json"
    assert json.loads(path.read_text(encoding="utf-8")) == document_contract()




def test_empty_pdf_requires_ocr(client):
    response = extract(client, pdf_bytes(blank=True))
    assert response.status_code == 200
    assert response.json()["status"] == "OCR_REQUIRED"
    assert "OCR_REQUIRED" in response.json()["warnings"]


def test_mixed_pdf_keeps_text_but_flags_image_only_page(client):
    from pypdf.generic import NumberObject
    writer = PdfWriter()
    writer.append(io.BytesIO(pdf_bytes()))
    page = writer.add_blank_page(width=100, height=100)
    image = DecodedStreamObject()
    image.set_data(b"\xff\xff\xff")
    image.update({NameObject("/Type"): NameObject("/XObject"), NameObject("/Subtype"): NameObject("/Image"),
                  NameObject("/Width"): NumberObject(1), NameObject("/Height"): NumberObject(1),
                  NameObject("/ColorSpace"): NameObject("/DeviceRGB"), NameObject("/BitsPerComponent"): NumberObject(8)})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/XObject"): DictionaryObject({NameObject("/Im0"): writer._add_object(image)})})
    content = DecodedStreamObject()
    content.set_data(b"q 100 0 0 100 0 0 cm /Im0 Do Q")
    page[NameObject("/Contents")] = writer._add_object(content)
    target = io.BytesIO()
    writer.write(target)
    result = extract(client, target.getvalue()).json()
    assert result["status"] == "OCR_REQUIRED"
    assert result["unit_count"] == 2
    assert result["spans"][0]["text"] == "Stack uses LIFO."
    assert result["spans"][1] == {"kind": "PAGE", "position": 2, "text": ""}


def test_mixed_pptx_flags_image_inside_group(client):
    from PIL import Image
    presentation = Presentation(io.BytesIO(pptx_bytes()))
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    image = io.BytesIO()
    Image.new("RGB", (10,10), "white").save(image, format="PNG")
    image.seek(0)
    group.shapes.add_picture(image, 0, 0)
    target = io.BytesIO()
    presentation.save(target)
    result = extract(client, target.getvalue(), "PPTX").json()
    assert result["status"] == "OCR_REQUIRED"
    assert result["unit_count"] == 2


def test_mixed_pptx_flags_image_inside_picture_placeholder(client):
    from PIL import Image
    from pptx.shapes.picture import Picture
    from pptx.shapes.placeholder import PicturePlaceholder
    presentation = Presentation(io.BytesIO(pptx_bytes()))
    slide = presentation.slides.add_slide(presentation.slide_layouts[8])
    placeholder = next(shape for shape in slide.placeholders if isinstance(shape, PicturePlaceholder))
    image = io.BytesIO()
    Image.new("RGB", (10, 10), "white").save(image, format="PNG")
    image.seek(0)
    picture = placeholder.insert_picture(image)
    assert isinstance(picture, Picture)
    assert picture.shape_type != 13
    target = io.BytesIO()
    presentation.save(target)
    result = extract(client, target.getvalue(), "PPTX").json()
    assert result["status"] == "OCR_REQUIRED"
    assert result["spans"][1]["text"] == ""


def test_strict_pdf_policy_rejects_recoverable_xref_damage(client):
    import re
    from pypdf import PdfReader
    damaged = re.sub(rb"startxref\s+\d+", b"startxref\n1", pdf_bytes())
    assert PdfReader(io.BytesIO(damaged), strict=False).pages[0].extract_text().strip() == "Stack uses LIFO."
    response = extract(client, damaged)
    assert response.status_code == 422
    assert response.json() == {"detail": "DOCUMENT_INVALID"}


def test_encrypted_pdf_is_rejected_without_password_or_source_leak(client):
    response = extract(client, pdf_bytes(encrypted=True))
    assert response.status_code == 422
    assert response.json() == {"detail": "DOCUMENT_ENCRYPTED"}


@pytest.mark.parametrize("notes", [True, False])
def test_pptx_slide_text_and_notes_opt_in(client, notes):
    response = extract(client, pptx_bytes(), "PPTX", notes)
    assert response.status_code == 200
    result = response.json()
    assert result["unit_count"] == 1
    assert "Ngăn xếp" in result["spans"][0]["text"]
    assert any(s["kind"] == "NOTES" for s in result["spans"]) is notes


@pytest.mark.parametrize("source,data", [("PDF", b"not a PDF"), ("PPTX", b"%PDF-1.7")])
def test_spoofed_type_is_rejected(client, source, data):
    assert extract(client, data, source).status_code == 422


def test_internal_auth_and_media_type(client):
    assert client.post("/internal/documents/extract?source_type=PDF", content=b"x").status_code == 401
    assert client.post("/internal/documents/extract?source_type=PDF", headers={"X-Internal-Token": TOKEN}, content=b"x").status_code == 415


def test_zip_path_traversal_and_entry_limit(client):
    for names in [["../escape.xml"], [f"item{i}" for i in range(5001)]]:
        target = io.BytesIO()
        with zipfile.ZipFile(target, "w") as archive:
            for name in names:
                archive.writestr(name, "x")
        assert extract(client, target.getvalue(), "PPTX").status_code == 422


def test_upload_actual_byte_limit(client):
    assert extract(client, b"x" * (20*1024*1024+1)).status_code == 413


def test_deadline_kills_child_and_releases_capacity(monkeypatch):
    original = asyncio.create_subprocess_exec
    children = []

    async def slow(*args, **kwargs):
        process = await original(sys.executable, "-c", "import time; time.sleep(10)", **kwargs)
        children.append(process)
        return process

    async def stream():
        yield pdf_bytes()

    async def run():
        parser = DocumentParser(timeout_seconds=0.2, max_concurrency=1)
        monkeypatch.setattr(asyncio, "create_subprocess_exec", slow)
        with pytest.raises(ExtractionError, match="DOCUMENT_TIMEOUT"):
            await parser.extract(stream(), "PDF", False)
        assert children[0].returncode is not None
        monkeypatch.setattr(asyncio, "create_subprocess_exec", original)
        parser.timeout_seconds = 10
        assert (await parser.extract(stream(), "PDF", False)).status == "EXTRACTED"
    asyncio.run(run())


def test_unsupported_subprocess_loop_falls_back_to_threaded_child(monkeypatch, caplog):
    caplog.set_level("INFO")
    original = subprocess.Popen
    environments = []
    async def unsupported(*args, **kwargs):
        raise NotImplementedError("PRIVATE loop details")
    def capture(*args, **kwargs):
        environments.append(kwargs["env"])
        return original(*args, **kwargs)
    monkeypatch.setattr(asyncio, "create_subprocess_exec", unsupported)
    monkeypatch.setattr(subprocess, "Popen", capture)
    monkeypatch.setenv("GEMINI_API_KEY", "synthetic-not-a-real-key")
    async def stream():
        yield pdf_bytes()
    async def run():
        parser = DocumentParser(max_concurrency=1)
        assert (await parser.extract(stream(), "PDF", False)).status == "EXTRACTED"
        assert parser._slots._value == 1
    asyncio.run(run())
    assert set(environments[0]) <= {"PYTHONUTF8", "SystemRoot", "WINDIR", "TEMP", "TMP"}
    assert "document_subprocess_thread_fallback" in caplog.text
    assert "PRIVATE" not in caplog.text


def test_threaded_child_is_killed_at_deadline(monkeypatch):
    original = subprocess.Popen
    children = []
    async def unsupported(*args, **kwargs):
        raise NotImplementedError
    def slow(args, **kwargs):
        children.append(original([sys.executable, "-c", "import time; time.sleep(10)"], **kwargs))
        return children[-1]
    monkeypatch.setattr(asyncio, "create_subprocess_exec", unsupported)
    monkeypatch.setattr(subprocess, "Popen", slow)
    async def stream():
        yield pdf_bytes()
    async def run():
        parser = DocumentParser(timeout_seconds=0.5, max_concurrency=1)
        with pytest.raises(ExtractionError, match="DOCUMENT_TIMEOUT"):
            await parser.extract(stream(), "PDF", False)
        assert children[0].poll() is not None
        assert parser._slots._value == 1
    asyncio.run(run())


@pytest.mark.skipif(sys.platform != "win32", reason="Windows selector loops lack subprocess support")
def test_windows_selector_loop_extracts_like_uvicorn_reload():
    async def stream():
        yield pdf_bytes()
    async def run():
        return await DocumentParser().extract(stream(), "PDF", False)
    assert asyncio.run(run(), loop_factory=asyncio.SelectorEventLoop).status == "EXTRACTED"


@pytest.mark.parametrize("stop", ["cancel", "deadline"])
@pytest.mark.parametrize("spawn_fails", [False, True])
def test_threaded_spawn_cancellation_retains_and_reaps_child(monkeypatch, caplog, stop, spawn_fails):
    original = subprocess.Popen
    main_thread = threading.get_ident()
    entered, release = threading.Event(), threading.Event()
    children = []
    folders = []

    async def unsupported(*args, **kwargs):
        raise NotImplementedError

    def delayed_spawn(args, **kwargs):
        assert threading.get_ident() != main_thread
        folders.append(Path(args[4]).parent)
        entered.set()
        if not release.wait(5):
            raise RuntimeError("Test did not release spawn")
        if spawn_fails:
            raise OSError("PRIVATE spawn details")
        child = original([sys.executable, "-c", "import time; time.sleep(10)"], **kwargs)
        children.append(child)
        return child

    monkeypatch.setattr(asyncio, "create_subprocess_exec", unsupported)
    monkeypatch.setattr(subprocess, "Popen", delayed_spawn)

    async def stream():
        yield pdf_bytes()

    async def run():
        parser = DocumentParser(timeout_seconds=0.05 if stop == "deadline" else 5, max_concurrency=1)
        active = asyncio.create_task(parser.extract(stream(), "PDF", False))
        try:
            assert await asyncio.to_thread(entered.wait, 2)
            if stop == "cancel":
                active.cancel()
                await asyncio.sleep(0.02)
            else:
                await asyncio.sleep(0.1)
            expected = asyncio.CancelledError if stop == "cancel" else ExtractionError
            with pytest.raises(expected) as error:
                await asyncio.wait_for(active, 0.5)
            if stop == "deadline":
                assert error.value.code == "DOCUMENT_TIMEOUT"
            # Request finishes before the spawning thread returns. Cleanup still
            # owns the files and slot, preventing unlimited abandoned spawns.
            assert not children and parser._cleanup_tasks
            assert folders[0].exists() and parser._slots._value == 0
            release.set()
            await asyncio.wait_for(asyncio.gather(*parser._cleanup_tasks, return_exceptions=True), 3)
            await asyncio.sleep(0)
            if spawn_fails:
                assert not children
                assert "document_spawn_cleanup_failed error_type=OSError" in caplog.text
                assert "PRIVATE" not in caplog.text
            else:
                assert len(children) == 1 and children[0].poll() is not None
            assert parser._slots._value == 1
            assert not parser._cleanup_tasks and not folders[0].exists()
            monkeypatch.setattr(subprocess, "Popen", original)
            parser.timeout_seconds = 10
            assert (await parser.extract(stream(), "PDF", False)).status == "EXTRACTED"
        finally:
            release.set()
            if not active.done():
                active.cancel()
            await asyncio.gather(active, return_exceptions=True)
            await asyncio.gather(*parser._cleanup_tasks, return_exceptions=True)

    asyncio.run(run())


def test_parser_child_receives_only_minimal_environment(monkeypatch):
    original = asyncio.create_subprocess_exec
    environments = []
    monkeypatch.setenv("GEMINI_API_KEY", "synthetic-not-a-real-key")
    monkeypatch.setenv("RAG_DATABASE_URL", "synthetic-not-a-real-url")
    monkeypatch.setenv("PYTHONPATH", "untrusted-module-path")
    async def capture(*args, **kwargs):
        environments.append(kwargs["env"])
        return await original(*args, **kwargs)
    monkeypatch.setattr(asyncio, "create_subprocess_exec", capture)
    async def stream():
        yield pdf_bytes()
    async def run():
        result = await DocumentParser().extract(stream(), "PDF", False)
        assert result.status == "EXTRACTED"
    asyncio.run(run())
    assert set(environments[0]) <= {"PYTHONUTF8", "SystemRoot", "WINDIR", "TEMP", "TMP"}
    assert environments[0]["PYTHONUTF8"] == "1"
